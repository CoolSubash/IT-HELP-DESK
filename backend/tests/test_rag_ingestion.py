"""
Ingestion pipeline tests (phase7.md #20's "Document ingestion: PDF, TXT,
Markdown", plus duplicate/versioning/failure handling from #19). Uses the
real `db` fixture (a real Postgres connection, see tests/conftest.py) and
the real DevEmbeddingProvider (app/rag/embeddings/dev_provider.py) except
where a test specifically needs to simulate a failing embedding call.
"""
import io
import pathlib
import uuid

import docx
import pytest
from reportlab.pdfgen import canvas

from app.errors import ConflictError, NotFoundError, ValidationError
from app.rag import ingestion_service
from app.rag.embeddings.provider import EmbeddingError


def _create_admin(db) -> uuid.UUID:
    admin_id = uuid.uuid4()
    with db.cursor() as cur:
        cur.execute(
            "INSERT INTO admins (id, name, email) VALUES (%s, %s, %s)",
            (admin_id, "Test Admin", f"{admin_id}@university.edu"),
        )
    return admin_id


def _pdf_bytes(text: str) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 720, text)
    pdf.save()
    return buffer.getvalue()


def _docx_bytes(heading: str, paragraph: str) -> bytes:
    buffer = io.BytesIO()
    document = docx.Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(paragraph)
    document.save(buffer)
    return buffer.getvalue()


def test_ingest_txt_creates_ready_document_with_chunks(db):
    admin_id = _create_admin(db)
    document = ingestion_service.ingest_document(
        db,
        title="Test TXT Guide",
        description="a test doc",
        category="OTHER",
        file_name="guide.txt",
        file_type="txt",
        file_bytes=b"Restart the VPN client, then try connecting again.",
        uploaded_by=admin_id,
    )
    assert document["status"] == "READY"
    assert document["error_message"] is None

    chunks, total = ingestion_service.list_chunks(db, document["id"], limit=20, offset=0)
    assert total == 1
    assert chunks[0]["chunk_index"] == 0
    assert "VPN client" in chunks[0]["content"]


def test_ingest_markdown_creates_ready_document_with_chunks(db):
    admin_id = _create_admin(db)
    markdown = "## VPN\n\nRestart the client.\n\n## WiFi\n\nForget and reconnect the network."
    document = ingestion_service.ingest_document(
        db,
        title="Test Markdown Guide",
        description=None,
        category="OTHER",
        file_name="guide.md",
        file_type="md",
        file_bytes=markdown.encode("utf-8"),
        uploaded_by=admin_id,
    )
    assert document["status"] == "READY"
    _, total = ingestion_service.list_chunks(db, document["id"], limit=20, offset=0)
    assert total >= 1


def test_ingest_pdf_creates_ready_document_with_chunks(db):
    admin_id = _create_admin(db)
    document = ingestion_service.ingest_document(
        db,
        title="Test PDF Guide",
        description=None,
        category="OTHER",
        file_name="guide.pdf",
        file_type="pdf",
        file_bytes=_pdf_bytes("VPN troubleshooting test content for ingestion."),
        uploaded_by=admin_id,
    )
    assert document["status"] == "READY"
    chunks, total = ingestion_service.list_chunks(db, document["id"], limit=20, offset=0)
    assert total == 1
    assert "VPN troubleshooting" in chunks[0]["content"]


def test_ingest_docx_creates_ready_document_with_chunks(db):
    admin_id = _create_admin(db)
    document = ingestion_service.ingest_document(
        db,
        title="Test DOCX Guide",
        description=None,
        category="OTHER",
        file_name="guide.docx",
        file_type="docx",
        file_bytes=_docx_bytes("VPN Guide", "Restart the VPN client after a password reset."),
        uploaded_by=admin_id,
    )
    assert document["status"] == "READY"
    chunks, total = ingestion_service.list_chunks(db, document["id"], limit=20, offset=0)
    assert total == 1
    assert "Restart the VPN client" in chunks[0]["content"]


def test_unsupported_file_type_raises_before_creating_any_row(db):
    admin_id = _create_admin(db)
    with pytest.raises(ValidationError):
        ingestion_service.ingest_document(
            db,
            title="Bad Type Doc",
            description=None,
            category=None,
            file_name="malware.exe",
            file_type="exe",
            file_bytes=b"whatever",
            uploaded_by=admin_id,
        )
    items, total = ingestion_service.list_documents(db, limit=20, offset=0, search="Bad Type Doc")
    assert total == 0


def test_empty_file_raises_before_creating_any_row(db):
    with pytest.raises(ValidationError):
        ingestion_service.ingest_document(
            db,
            title="Empty Doc",
            description=None,
            category=None,
            file_name="empty.txt",
            file_type="txt",
            file_bytes=b"",
        )
    items, total = ingestion_service.list_documents(db, limit=20, offset=0, search="Empty Doc")
    assert total == 0


def test_duplicate_title_without_replace_raises_conflict(db):
    ingestion_service.ingest_document(
        db, title="Duplicate Title Doc", description=None, category=None,
        file_name="a.txt", file_type="txt", file_bytes=b"first version content here",
    )
    with pytest.raises(ConflictError):
        ingestion_service.ingest_document(
            db, title="Duplicate Title Doc", description=None, category=None,
            file_name="b.txt", file_type="txt", file_bytes=b"second version content here",
        )


def test_replace_document_id_creates_new_version_and_archives_old(db):
    original = ingestion_service.ingest_document(
        db, title="Versioned Doc", description=None, category=None,
        file_name="v1.txt", file_type="txt", file_bytes=b"version one content here",
    )
    assert original["version"] == 1

    new_version = ingestion_service.ingest_document(
        db, title="Versioned Doc", description=None, category=None,
        file_name="v2.txt", file_type="txt", file_bytes=b"version two content here",
        replace_document_id=original["id"],
    )
    assert new_version["version"] == 2
    assert new_version["previous_version_id"] == original["id"]
    assert new_version["status"] == "READY"

    archived_original = ingestion_service.get_document(db, original["id"])
    assert archived_original["status"] == "ARCHIVED"


def test_replace_document_id_for_missing_document_raises_not_found(db):
    with pytest.raises(NotFoundError):
        ingestion_service.ingest_document(
            db, title="Anything", description=None, category=None,
            file_name="a.txt", file_type="txt", file_bytes=b"content",
            replace_document_id=uuid.uuid4(),
        )


def test_embedding_failure_leaves_document_failed_with_error_message(db, monkeypatch):
    class _FailingEmbeddingService:
        def generate_embeddings(self, texts):
            raise EmbeddingError("simulated embedding provider outage")

    monkeypatch.setattr(
        ingestion_service, "get_embedding_service", lambda: _FailingEmbeddingService()
    )

    document = ingestion_service.ingest_document(
        db, title="Doomed Doc", description=None, category=None,
        file_name="a.txt", file_type="txt", file_bytes=b"some content that will fail to embed",
    )
    assert document["status"] == "FAILED"
    assert "simulated embedding provider outage" in document["error_message"]

    _, total = ingestion_service.list_chunks(db, document["id"], limit=20, offset=0)
    assert total == 0


def test_connection_stays_usable_after_a_failed_ingestion(db, monkeypatch):
    """The SAVEPOINT/ROLLBACK TO SAVEPOINT in ingestion_service.py must
    leave `db` in a working state for whatever runs next in the same
    transaction -- this is exactly the scenario a pooled connection
    (app/database.py's get_db()) hits across two requests sharing a
    connection."""
    class _FailingEmbeddingService:
        def generate_embeddings(self, texts):
            raise EmbeddingError("simulated outage")

    monkeypatch.setattr(
        ingestion_service, "get_embedding_service", lambda: _FailingEmbeddingService()
    )
    failed = ingestion_service.ingest_document(
        db, title="Doomed Doc 2", description=None, category=None,
        file_name="a.txt", file_type="txt", file_bytes=b"content",
    )
    assert failed["status"] == "FAILED"

    monkeypatch.undo()  # restore the real (dev) embedding service
    succeeded = ingestion_service.ingest_document(
        db, title="Recovers Fine Doc", description=None, category=None,
        file_name="b.txt", file_type="txt", file_bytes=b"content that should embed fine",
    )
    assert succeeded["status"] == "READY"


def test_get_document_raises_not_found_for_missing_id(db):
    with pytest.raises(NotFoundError):
        ingestion_service.get_document(db, uuid.uuid4())


def test_list_documents_filters_by_status_and_category(db):
    ingestion_service.ingest_document(
        db, title="Filter Test VPN Doc", description=None, category="VPN",
        file_name="a.txt", file_type="txt", file_bytes=b"vpn content here",
    )
    items, total = ingestion_service.list_documents(db, limit=20, offset=0, category="VPN")
    assert total >= 1
    assert all(item["category"] == "VPN" for item in items)

    items, total = ingestion_service.list_documents(db, limit=20, offset=0, status="ARCHIVED", search="Filter Test VPN Doc")
    assert total == 0


def test_archive_document_sets_status_archived(db):
    document = ingestion_service.ingest_document(
        db, title="To Be Archived Doc", description=None, category=None,
        file_name="a.txt", file_type="txt", file_bytes=b"content to archive",
    )
    archived = ingestion_service.archive_document(db, document["id"])
    assert archived["status"] == "ARCHIVED"


def test_delete_document_hard_deletes_row_and_cascades_chunks(db):
    document = ingestion_service.ingest_document(
        db, title="To Be Deleted Doc", description=None, category=None,
        file_name="a.txt", file_type="txt", file_bytes=b"content to delete",
    )
    stored_path = pathlib.Path(document["storage_location"])
    assert stored_path.exists()  # sanity check before delete -- ingest_document() really wrote it

    ingestion_service.delete_document(db, document["id"])

    with pytest.raises(NotFoundError):
        ingestion_service.get_document(db, document["id"])

    with db.cursor() as cur:
        cur.execute("SELECT count(*) FROM knowledge_chunks WHERE document_id = %s", (document["id"],))
        assert cur.fetchone()[0] == 0

    # The underlying file must be gone too, not just the database row --
    # see ingestion_service.delete_document()'s docstring.
    assert not stored_path.exists()


def test_delete_document_is_idempotent_if_called_twice(db):
    """A second delete on an id that's already gone should behave like
    any other missing-document call (NotFoundError from get_document()),
    not raise a different, confusing error from the storage layer."""
    document = ingestion_service.ingest_document(
        db, title="Double Delete Doc", description=None, category=None,
        file_name="a.txt", file_type="txt", file_bytes=b"content",
    )
    ingestion_service.delete_document(db, document["id"])
    with pytest.raises(NotFoundError):
        ingestion_service.delete_document(db, document["id"])
